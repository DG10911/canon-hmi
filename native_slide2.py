#!/usr/bin/env python3
"""Replace Slide 2 flat image with fully editable native PowerPoint shapes."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

F = "/Users/devanshgoenka/conductor/workspaces/dgrfrfeerf/santo-domingo/.context/EcoDrive_Intelligence_Round1_TEMPLATE.pptx"
GREEN=RGBColor(0x3D,0xCD,0x58); DARK=RGBColor(0x1A,0x1A,0x1A); GREY=RGBColor(0x4D,0x4D,0x4D)
LGREY=RGBColor(0xE9,0xEC,0xEF); MGREY=RGBColor(0xCE,0xD4,0xDA); WHITE=RGBColor(0xFF,0xFF,0xFF); BLUE=RGBColor(0x2B,0x6C,0xB0)

prs=Presentation(F)
s=prs.slides[1]  # slide 2

# 1) remove the existing flat picture
for sh in list(s.shapes):
    if sh.shape_type==13:  # PICTURE
        sh._element.getparent().remove(sh._element)

def rect(x,y,w,h,text="",fill=WHITE,line=GREEN,tc=DARK,fs=10,bold=True,lw=1.5,
         align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE,shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp=s.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb=fill
    sp.line.color.rgb=line; sp.line.width=Pt(lw)
    sp.shadow.inherit=False
    tf=sp.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    for m in ("margin_left","margin_right","margin_top","margin_bottom"): setattr(tf,m,Inches(0.03))
    lines=text.split("\n")
    for i,ln in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; r=p.add_run(); r.text=ln
        r.font.size=Pt(fs); r.font.bold=bold; r.font.color.rgb=tc
    return sp

def label(x,y,text,fs=8,tc=GREY,bold=True):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(2.2),Inches(0.25))
    tf=tb.text_frame; tf.margin_left=0; tf.margin_top=0
    p=tf.paragraphs[0]; r=p.add_run(); r.text=text
    r.font.size=Pt(fs); r.font.bold=bold; r.font.color.rgb=tc
    return tb

def conn(x1,y1,x2,y2,color=GREY,lw=1.75,label_txt=None):
    c=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    c.line.color.rgb=color; c.line.width=Pt(lw)
    ln=c.line._get_or_add_ln()
    te=ln.makeelement(qn('a:tailEnd'),{'type':'triangle','w':'med','len':'med'}); ln.append(te)
    if label_txt:
        tb=s.shapes.add_textbox(Inches((x1+x2)/2-0.9),Inches(min(y1,y2)-0.02),Inches(1.8),Inches(0.22))
        tf=tb.text_frame; tf.margin_left=0; tf.margin_top=0
        p=tf.paragraphs[0]; p.alignment=PP_ALIGN.CENTER; r=p.add_run(); r.text=label_txt
        r.font.size=Pt(8); r.font.bold=True; r.font.color.rgb=BLUE
    return c

# ---- layer backgrounds (left) ----
LX,LWD=0.3,7.65
layers=[("LAYER – CLOUD",1.42,0.82),("LAYER – EDGE",2.34,0.72),
        ("LAYER – CONTROL",3.16,1.02),("LAYER – FIELD",4.28,1.92)]
for name,y,h in layers:
    bg=rect(LX,y,LWD,h,"",fill=LGREY,line=MGREY,lw=1.0,shape=MSO_SHAPE.RECTANGLE)
    label(LX+0.08,y+0.04,name,fs=8,tc=GREY)

# ---- device boxes ----
rect(2.1,1.60,3.9,0.5,"EcoStruxure Machine Advisor\nRemote dashboards • OEE • alerts",line=BLUE,fs=9)   # cloud
rect(2.1,2.48,3.9,0.46,"Edge Gateway (Harmony IPC)\nLocal historian + analytics",fs=9)                    # edge
m241=rect(0.95,3.42,2.7,0.58,"Modicon M241\nPLC (IEC 61131-3)",fs=9)                                      # control
hmi =rect(4.05,3.42,2.5,0.58,"Harmony GTU\nHMI (ISA-101)",fs=9)
rect(0.5,5.28,2.3,0.72,"Pressure / Flow /\nLevel Transmitters",fs=8.5)                                    # field
rect(3.0,5.28,2.3,0.72,"ATV630 VFD\n↔ Motor + Pump",fs=8.5)
rect(5.5,5.28,2.3,0.72,"PowerLogic PM5300\nEnergy Meter",fs=8.5)

# ---- arrows (upward through the stack) + PLC<->HMI ----
conn(2.0,5.28,2.3,4.02,label_txt="Modbus TCP / 4–20 mA")   # field -> control
conn(4.0,3.42,4.0,2.96,label_txt="Modbus TCP")             # control -> edge
conn(4.0,2.48,4.0,2.12,color=BLUE,label_txt="MQTT / TLS")  # edge -> cloud
conn(3.65,3.71,4.05,3.71,color=GREEN)                      # PLC -> HMI

# ---- approach column (right) ----
label(8.25,1.42,"Approach",fs=15,tc=DARK)
und=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(8.25),Inches(1.78),Inches(1.0),Inches(0.035))
und.fill.solid(); und.fill.fore_color.rgb=GREEN; und.line.fill.background(); und.shadow.inherit=False
tb=s.shapes.add_textbox(Inches(8.25),Inches(1.95),Inches(4.85),Inches(4.1))
tf=tb.text_frame; tf.word_wrap=True
bullets=[
 "Use case: friction-dominated water pumping station (30 kW).",
 "VFD speed control replaces valve throttling → cube-law savings (P ∝ N³).",
 "Savings calibrated to the real system curve (static vs. friction head).",
 "Closed loop: SENSE → CONTROL → VISUALIZE → ANALYZE → OPTIMIZE → ACT → LEARN.",
 "Modbus TCP common layer — native to all four Schneider devices.",
]
for i,b in enumerate(bullets):
    p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
    p.space_after=Pt(10)
    r=p.add_run(); r.text="▶  "+b; r.font.size=Pt(11.5); r.font.color.rgb=GREY
    r.font.bold=False

# ---- KPI strip (bottom) ----
pills=["~46%  Energy ↓","SEC 0.58→0.31 kWh/m³","₹4–6 L / year saved","Payback < 2 years"]
pw=3.02; gap=0.15; x0=0.3
for i,p in enumerate(pills):
    rect(x0+i*(pw+gap),6.42,pw,0.5,p,fill=GREEN,line=GREEN,tc=WHITE,fs=11.5)

prs.save(F)
print("Slide 2 rebuilt with editable native shapes.")
